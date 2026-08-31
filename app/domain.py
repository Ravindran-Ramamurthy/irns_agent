import io
import logging
from dataclasses import dataclass
from typing import Any, Dict, List

from docx import Document as DocxDocument

from app import settings
from app.server_client import ServerClient

log = logging.getLogger(__name__)

_SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx"}

# The 3 lookup tables this agent needs to turn an LLM decision's human-readable
# codes into the FK ids ActionEvent/RiskEvent actually store. See
# ai.smartgateway.backend.irns.master.model.IrnsMasterType on the Java side -
# ALERT_CODE/EVENT_TYPE/RESPONSE_CODE are irrelevant here, this agent only
# ever creates action events and updates a risk event's status.
_MASTER_TYPES = ("ACTION_TYPE", "ACTION_STATUS", "STATUS")


class DomainError(Exception):
    """Anything wrong with the IRNS knowledge domain itself - missing/duplicate
    domain, no supported file, or nothing extractable from it. Distinct from
    ServerError (network/HTTP failure talking to smartgateway) so callers can
    tell "smartgateway is down" apart from "the tenant hasn't set the IRNS
    domain up correctly"."""


@dataclass(frozen=True)
class IngestResult:
    chunks: List[str]
    masters: Dict[str, Dict[str, Dict[str, Any]]]  # master type -> code -> row


async def ingest(client: ServerClient) -> IngestResult:
    """Reads every supported file in the tenant's "IRNS" knowledge domain,
    chunks their text, and fetches the master lookup tables - everything
    /init keeps in app.store's static, in-memory collection. Raises
    DomainError for any setup problem, ServerError if smartgateway itself
    couldn't be reached."""
    domain = await _find_domain(client)
    files = await _find_files(client, domain)

    all_chunks: List[str] = []
    for file in files:
        raw = await client.download_file(file["id"])
        text = _extract_text(raw, file.get("fileName") or "")
        all_chunks.extend(_chunk(text))

    if not all_chunks:
        raise DomainError(
            f"No text could be extracted from the '{settings.KNOWLEDGE_DOMAIN_NAME}' domain's file(s)")

    masters: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for master_type in _MASTER_TYPES:
        rows = await client.get_master_all(master_type)
        masters[master_type] = {row["code"]: row for row in rows if row.get("code")}

    return IngestResult(chunks=all_chunks, masters=masters)


async def _find_domain(client: ServerClient) -> Dict[str, Any]:
    domains = await client.list_domains()
    matches = [d for d in domains if (d.get("name") or "") == settings.KNOWLEDGE_DOMAIN_NAME]

    if not matches:
        raise DomainError(f"No knowledge domain named '{settings.KNOWLEDGE_DOMAIN_NAME}' was found")
    if len(matches) > 1:
        raise DomainError(
            f"Found {len(matches)} knowledge domains named '{settings.KNOWLEDGE_DOMAIN_NAME}' - there must be exactly one")
    return matches[0]


async def _find_files(client: ServerClient, domain: Dict[str, Any]) -> List[Dict[str, Any]]:
    files = await client.list_domain_files(domain["id"])
    if not files:
        raise DomainError(f"The '{settings.KNOWLEDGE_DOMAIN_NAME}' knowledge domain has no file")

    supported = [f for f in files if _suffix(f.get("fileName") or "") in _SUPPORTED_EXTENSIONS]
    if not supported:
        raise DomainError(
            f"The '{settings.KNOWLEDGE_DOMAIN_NAME}' knowledge domain has no .txt/.md/.docx file")
    return supported


def _suffix(file_name: str) -> str:
    dot = file_name.rfind(".")
    return file_name[dot:].lower() if dot >= 0 else ""


def _extract_text(raw: bytes, file_name: str) -> str:
    if _suffix(file_name) == ".docx":
        doc = DocxDocument(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1", errors="replace")


def _chunk(text: str) -> List[str]:
    """Naive fixed-size character splitter with overlap - no embeddings, no
    retrieval ranking. Good enough for a single SRS-sized document; a real
    vector store would be the next step if the document set grows."""
    text = text.strip()
    if not text:
        return []

    size = settings.CHUNK_SIZE_CHARS
    step = max(size - settings.CHUNK_OVERLAP_CHARS, 1)

    chunks: List[str] = []
    start = 0
    while start < len(text):
        chunks.append(text[start:start + size])
        start += step
    return chunks
